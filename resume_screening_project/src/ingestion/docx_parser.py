import logging
import docx

logger = logging.getLogger(__name__)


def extract_text_from_docx(docx_path: str) -> str:
    try:
        document = docx.Document(docx_path)
        text_parts = []
        for para in document.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        full_text = "\n".join(text_parts)
        if not full_text.strip():
            logger.warning(f"Empty text extracted from {docx_path}")
        return full_text
    except Exception as e:
        logger.error(f"Failed to parse DOCX {docx_path}: {e}")
        raise
